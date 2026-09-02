import openpyxl

from app.models import InvoiceData, LineItem
from app.excel_generator import generate_excel


def _sample_invoice():
    return InvoiceData(
        from_name="SLV Cleaning Services",
        from_phone="508-901-1384",
        from_address="69, Cato lane",
        bill_to_name="Kristin Paton",
        bill_to_phone="555-1234",
        bill_to_address="1 Main St",
        invoice_no="3",
        invoice_date="30/05",
        date_due="10/06",
        payment_method="Cash",
        hourly_rate=35.0,
        items=[
            LineItem(date="01/04", client="11 pleasant", hours=4.0),
            LineItem(date="02/04", client="Caren", hours=1.0),
        ],
    )


def _find_cell_with_value(ws, value):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == value:
                return cell
    return None


def test_generate_excel_writes_bill_to_and_from_names(tmp_path):
    out_path = tmp_path / "invoice.xlsx"
    generate_excel(_sample_invoice(), out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    assert _find_cell_with_value(ws, "Kristin Paton") is not None
    assert _find_cell_with_value(ws, "SLV Cleaning Services") is not None


def test_generate_excel_writes_item_rows_and_total(tmp_path):
    out_path = tmp_path / "invoice.xlsx"
    generate_excel(_sample_invoice(), out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    assert _find_cell_with_value(ws, "11 pleasant") is not None
    assert _find_cell_with_value(ws, "01/04") is not None
    assert _find_cell_with_value(ws, 4.0) is not None
    total_cell = _find_cell_with_value(ws, 175.0)
    assert total_cell is not None


def test_generate_excel_writes_hourly_rate_in_meta_block(tmp_path):
    out_path = tmp_path / "invoice.xlsx"
    generate_excel(_sample_invoice(), out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    assert _find_cell_with_value(ws, 35.0) is not None


def test_generate_excel_item_headers_are_date_client_amount(tmp_path):
    out_path = tmp_path / "invoice.xlsx"
    generate_excel(_sample_invoice(), out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    assert _find_cell_with_value(ws, "Date") is not None
    assert _find_cell_with_value(ws, "Client") is not None
    assert _find_cell_with_value(ws, "Amount") is not None


def test_generate_excel_header_cells_use_slv_blue_fill(tmp_path):
    out_path = tmp_path / "invoice.xlsx"
    generate_excel(_sample_invoice(), out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    cell = _find_cell_with_value(ws, "Invoice no.")
    assert cell is not None
    assert cell.fill.fgColor.rgb == "FF6D9EEB"
